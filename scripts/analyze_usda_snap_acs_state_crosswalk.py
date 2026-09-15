#!/usr/bin/env python3
"""Crosswalk FY2025 USDA SNAP state rates with existing 2024 ACS state context."""

from __future__ import annotations

import argparse
import json
import math
from pathlib import Path

from openpyxl import load_workbook


STATE_NAMES = {
    "01": "Alabama", "02": "Alaska", "04": "Arizona", "05": "Arkansas",
    "06": "California", "08": "Colorado", "09": "Connecticut", "10": "Delaware",
    "11": "District of Columbia", "12": "Florida", "13": "Georgia", "15": "Hawaii",
    "16": "Idaho", "17": "Illinois", "18": "Indiana", "19": "Iowa", "20": "Kansas",
    "21": "Kentucky", "22": "Louisiana", "23": "Maine", "24": "Maryland",
    "25": "Massachusetts", "26": "Michigan", "27": "Minnesota", "28": "Mississippi",
    "29": "Missouri", "30": "Montana", "31": "Nebraska", "32": "Nevada",
    "33": "New Hampshire", "34": "New Jersey", "35": "New Mexico", "36": "New York",
    "37": "North Carolina", "38": "North Dakota", "39": "Ohio", "40": "Oklahoma",
    "41": "Oregon", "42": "Pennsylvania", "44": "Rhode Island", "45": "South Carolina",
    "46": "South Dakota", "47": "Tennessee", "48": "Texas", "49": "Utah",
    "50": "Vermont", "51": "Virginia", "53": "Washington", "54": "West Virginia",
    "55": "Wisconsin", "56": "Wyoming",
}


def read_snap_rates(path: Path) -> dict[str, float]:
    sheet = load_workbook(path, data_only=True, read_only=True).active
    return {
        str(row[0].value): float(row[1].value)
        for row in sheet.iter_rows(min_row=3, max_col=2)
        if row[0].value and isinstance(row[1].value, (int, float))
    }


def pearson(rows: list[dict], left: str, right: str) -> float:
    xs = [float(row[left]) for row in rows]
    ys = [float(row[right]) for row in rows]
    mean_x, mean_y = sum(xs) / len(xs), sum(ys) / len(ys)
    numerator = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    denominator = math.sqrt(sum((x - mean_x) ** 2 for x in xs) * sum((y - mean_y) ** 2 for y in ys))
    return numerator / denominator if denominator else None


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--snap-xlsx", type=Path, required=True)
    parser.add_argument("--acs-state-json", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()

    snap = read_snap_rates(args.snap_xlsx)
    acs = json.loads(args.acs_state_json.read_text())
    rows = []
    for item in acs["states"]:
        state = STATE_NAMES.get(item["geo_id"][-2:])
        if state not in snap:
            continue
        rows.append({
            "state": state,
            "geo_id": item["geo_id"],
            "snap_participation_rate": snap[state],
            "median_household_income": item["median_household_income"],
            "zero_vehicle_households": item["zero_vehicle"],
            "households": item["households"],
            "long_commute_workers": item["long_commute"],
            "workers": item["workers"],
            "zero_vehicle_rate": 100 * item["zero_vehicle"] / item["households"],
            "long_commute_rate": 100 * item["long_commute"] / item["workers"],
        })
    rows.sort(key=lambda row: (row["median_household_income"], row["state"]))
    for index, row in enumerate(rows):
        row["acs_income_quartile"] = min(4, index * 4 // len(rows) + 1)

    quartiles = {}
    for quartile in range(1, 5):
        group = [row for row in rows if row["acs_income_quartile"] == quartile]
        households = sum(row["households"] for row in group)
        workers = sum(row["workers"] for row in group)
        quartiles[str(quartile)] = {
            "states": len(group),
            "median_income_range": [min(row["median_household_income"] for row in group), max(row["median_household_income"] for row in group)],
            "mean_snap_participation_rate": round(sum(row["snap_participation_rate"] for row in group) / len(group), 2),
            "zero_vehicle_households_share": round(100 * sum(row["zero_vehicle_households"] for row in group) / households, 2),
            "long_commute_share": round(100 * sum(row["long_commute_workers"] for row in group) / workers, 2),
        }

    result = {
        "format": "usda-snap-acs-state-crosswalk-v1",
        "source": "USDA FY2025 state participation chart data joined to existing 2024 ACS state context",
        "state_count": len(rows),
        "quartile_definition": "51 matched states/DC sorted by ACS B19013 median household income and split into four near-equal groups; ACS mobility measures are pooled by their table totals, while SNAP participation is summarized as the unweighted mean of state rates.",
        "correlations": {
            "snap_vs_median_household_income": round(pearson(rows, "snap_participation_rate", "median_household_income"), 3),
            "snap_vs_zero_vehicle_rate": round(pearson(rows, "snap_participation_rate", "zero_vehicle_rate"), 3),
            "snap_vs_long_commute_rate": round(pearson(rows, "snap_participation_rate", "long_commute_rate"), 3),
        },
        "quartiles": quartiles,
        "states": rows,
    }
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()

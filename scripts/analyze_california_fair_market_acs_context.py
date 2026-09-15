#!/usr/bin/env python3
"""Join the California FAIR/private-market ZIP screen to ACS context."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader
import re

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "analysis/projects/us-housing-insurance-risk/data"
OUT = DATA / "california-fair-market-acs-context-2022-join.json"


def fair_rows() -> dict[str, float]:
    text = "\n".join(page.extract_text() or "" for page in PdfReader(DATA / "california-fair-plan-vs-voluntary-2022.pdf").pages)
    pat = re.compile(r"^(.*?)\s+(\d{5})\s+(.*?)\s+([\d,]+|-|—)\s+([\d,]+|-|—)\s+([\d.]+)%\s*$")
    rows = {}
    for line in text.splitlines():
        m = pat.match(line)
        if m:
            rows[m.group(2)] = float(m.group(6))
    return rows


def summarize(rows: list[dict], label: str) -> dict:
    renewed = sum(r["renewed"] for r in rows)
    nonrenewed = sum(r["nonrenewed"] for r in rows)
    return {"group": label, "zip_rows": len(rows), "renewed": renewed, "nonrenewed": nonrenewed, "nonrenewal_rate": round(nonrenewed / (renewed + nonrenewed), 6) if renewed + nonrenewed else None, "mean_fair_share": round(sum(r["fair_share"] for r in rows) / len(rows), 3) if rows else None}


def main() -> int:
    income = {}
    with (DATA / "acs2022-b19013.dat").open() as fh:
        headers = next(fh).rstrip("\n").split("|")
        for line in fh:
            parts = line.rstrip("\n").split("|")
            item = dict(zip(headers, parts))
            geo = item.get("GEO_ID", "")
            if geo.startswith("860Z200US"):
                try:
                    income[geo[-5:]] = int(item["B19013_E001"])
                except (TypeError, ValueError):
                    pass
    acs = {}
    with (DATA / "acs2022-b25024.dat").open() as fh:
        headers = next(fh).rstrip("\n").split("|")
        for line in fh:
            parts = line.rstrip("\n").split("|")
            item = dict(zip(headers, parts))
            geo = item.get("GEO_ID", "")
            z = geo[-5:] if geo.startswith("860Z200US") else None
            try:
                total = int(item["B25024_E001"])
                detached = int(item["B25024_E002"])
                mobile = int(item["B25024_E010"])
            except (TypeError, ValueError):
                continue
            if z and z in income and income[z] >= 0 and total > 0:
                acs[z] = {"median_household_income": income[z], "structure_units": total, "detached_share": detached / total * 100, "mobile_home_share": mobile / total * 100}

    fair = fair_rows()
    ws = load_workbook(DATA / "california-residential-insurance-zip-2020-2023.xlsx", read_only=True, data_only=True).active
    joined = []
    for _county, z, year, _new, renewed, nonrenewed in ws.iter_rows(min_row=2, values_only=True):
        z = str(z).zfill(5)
        if year == 2022 and z in fair and z in acs:
            joined.append({"zip": z, "fair_share": fair[z], "renewed": renewed or 0, "nonrenewed": nonrenewed or 0, **acs[z]})

    joined.sort(key=lambda r: r["median_household_income"])
    q = len(joined) // 4
    income_groups = [summarize(joined[:q], "lowest_income_quartile"), summarize(joined[-q:], "highest_income_quartile")]
    by_detached = sorted(joined, key=lambda r: r["detached_share"])
    detached_groups = [summarize(by_detached[:q], "lowest_detached_share_quartile"), summarize(by_detached[-q:], "highest_detached_share_quartile")]
    result = {"format": "california-fair-market-acs-context-2022-v1", "period": "2022 FAIR share, voluntary decisions, and ACS 2022 5-year ZCTA context", "joined_zip_rows": len(joined), "fair_zip_rows": len(fair), "acs_zcta_rows": len(acs), "join_rate_against_fair_rows": round(len(joined) / len(fair), 6) if fair else None, "income_quartiles": income_groups, "structure_quartiles": detached_groups, "classification": "ZIP-level descriptive context join; ACS ZCTAs approximate ZIP geography and FAIR/voluntary denominators differ.", "limits": ["ACS is a five-year place estimate, not household insurance income or property-level composition.", "The voluntary workbook's counts and the FAIR table's residential-structure share have different universes.", "Equal-count quartiles are unweighted by structures, policies, or population.", "This does not identify individual transitions, reasons for nonrenewal, coverage adequacy, repairs, or moves."], "joined_rows": joined}
    OUT.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({k: result[k] for k in ("joined_zip_rows", "fair_zip_rows", "acs_zcta_rows", "join_rate_against_fair_rows", "income_quartiles", "structure_quartiles")}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

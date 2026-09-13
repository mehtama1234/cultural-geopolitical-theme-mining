#!/usr/bin/env python3
"""Compare Census county BFS application flows with CBP establishment stocks.

This is a scale diagnostic only.  It deliberately does not call the ratio a
formation, survival, or conversion rate.
"""

from __future__ import annotations

import argparse
import csv
import io
import statistics
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


def read_bfs(path: Path) -> dict[tuple[str, str], tuple[float, float, float]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook.active
    values: dict[tuple[str, str], tuple[float, float, float]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        try:
            key = (str(row[3]).zfill(2), str(row[4]).zfill(3))
            years = tuple(float(str(row[index]).replace(",", "")) for index in (23, 24, 25))
        except (TypeError, ValueError, IndexError):
            continue
        values[key] = years  # 2023, 2024, 2025
    return values


def read_cbp(path: Path) -> dict[tuple[str, str], int]:
    values: dict[tuple[str, str], int] = {}
    with ZipFile(path) as archive:
        members = [name for name in archive.namelist() if name.endswith(".txt")]
        if len(members) != 1:
            raise ValueError(f"expected one CBP text member, found {members}")
        with archive.open(members[0]) as raw:
            reader = csv.DictReader(io.TextIOWrapper(raw, encoding="latin1"))
            for row in reader:
                if row.get("naics") != "------":
                    continue
                try:
                    values[(row["fipstate"], row["fipscty"])] = int(row["est"])
                except (KeyError, TypeError, ValueError):
                    continue
    return values


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--bfs", type=Path, required=True)
    parser.add_argument("--cbp", type=Path, required=True)
    parser.add_argument("--minimum-establishments", type=int, default=10)
    args = parser.parse_args()

    bfs = read_bfs(args.bfs)
    cbp = read_cbp(args.cbp)
    matched = set(bfs) & set(cbp)
    usable = sorted(key for key in matched if cbp[key] >= args.minimum_establishments)

    print(f"bfs_counties={len(bfs)}")
    print(f"cbp_counties={len(cbp)}")
    print(f"matched_counties={len(matched)}")
    print(f"usable_counties={len(usable)}")
    for year, index in ((2023, 0), (2024, 1), (2025, 2)):
        ratios = [bfs[key][index] / cbp[key] for key in usable]
        print(f"median_ba{year}_over_cbp_est={statistics.median(ratios):.3f}")

    # Show the largest ratios for manual QA.  These are not findings.
    ranked = sorted(
        ((bfs[key][1] / cbp[key], key, bfs[key][1], cbp[key]) for key in usable),
        reverse=True,
    )
    for ratio, (state, county), applications, establishments in ranked[:5]:
        print(
            "qa_high_2024="
            f"{state}-{county},applications={applications:g},"
            f"establishments={establishments},ratio={ratio:.3f}"
        )


if __name__ == "__main__":
    main()

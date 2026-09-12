#!/usr/bin/env python3
"""Audit the structure of Treasury FIO's supporting metrics workbook.

This intentionally performs a structural audit only. It does not publish or
interpret ZIP-level rows, and it does not replace Treasury's aggregate analysis.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import openpyxl


EXPECTED_HEADERS = [
    "ZIP Code",
    "Year",
    "Policy Decile Grouping",
    "Claim Frequency",
    "Claim Severity",
    "Loss Ratio",
    "Premiums Per Policy",
    "Nonrenewal Rate",
    "Nonpayment Cancellation Rate",
    "Other than Nonpayment Cancellation Rate",
]


def audit(path: Path) -> dict:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for worksheet in workbook.worksheets:
        sheets.append(
            {
                "name": worksheet.title,
                "rows": worksheet.max_row,
                "columns": worksheet.max_column,
            }
        )

    metrics = workbook["Supporting Underlying Metrics"]
    headers = [cell.value for cell in next(metrics.iter_rows(min_row=1, max_row=1))]
    if headers != EXPECTED_HEADERS:
        raise ValueError(f"unexpected metrics headers: {headers!r}")
    return {
        "format": "us-housing-insurance-fio-workbook-audit-v1",
        "source_file": path.name,
        "sheets": sheets,
        "metrics_headers": headers,
        "metrics_data_rows": metrics.max_row - 1,
        "structural_audit_only": True,
        "raw_rows_published": False,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path)
    args = parser.parse_args()
    result = audit(args.workbook)
    rendered = json.dumps(result, indent=2) + "\n"
    if args.output:
        args.output.write_text(rendered)
    else:
        print(rendered, end="")


if __name__ == "__main__":
    main()

#!/usr/bin/env python3
"""Audit coverage and field completeness in Treasury's public FIO workbook."""

from __future__ import annotations

import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
INPUT = ROOT / "analysis/projects/us-housing-insurance-risk/data/treasury-fio-homeowners-insurance-2018-2022.xlsx"
OUTPUT = ROOT / "analysis/projects/us-housing-insurance-risk/data/treasury-fio-workbook-audit.json"


def main() -> int:
    ws = load_workbook(INPUT, read_only=True, data_only=True)["Supporting Underlying Metrics"]
    headers = [str(x) for x in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]
    years = Counter()
    zips_by_year = defaultdict(set)
    deciles_by_year = defaultdict(Counter)
    nonnull = Counter()
    rows = 0
    skipped = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in row):
            continue
        if row[0] is None or row[1] is None:
            skipped += 1
            continue
        rows += 1
        year = int(row[1])
        years[str(year)] += 1
        zips_by_year[str(year)].add(str(int(row[0])).zfill(5))
        if row[2] is not None:
            deciles_by_year[str(year)][str(int(row[2]))] += 1
        for header, value in zip(headers, row):
            if value is not None:
                nonnull[header] += 1
    result = {
        "format": "treasury-fio-workbook-audit-v1",
        "input": "treasury-fio-homeowners-insurance-2018-2022.xlsx",
        "sheet": "Supporting Underlying Metrics",
        "headers": headers,
        "data_rows": rows,
        "skipped_rows_missing_zip_or_year": skipped,
        "years": dict(sorted(years.items())),
        "unique_zip_codes_by_year": {year: len(zips) for year, zips in sorted(zips_by_year.items())},
        "policy_count_decile_rows_by_year": {year: dict(sorted(counts.items())) for year, counts in sorted(deciles_by_year.items())},
        "nonnull_cells_by_field": dict(nonnull),
        "interpretive_note": "Policy Decile Grouping is a relative grouping by number of policies in a ZIP code. It is not the climate-risk quintile used in Treasury's headline premium and nonrenewal comparisons.",
    }
    OUTPUT.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"data_rows": rows, "years": dict(sorted(years.items())), "output": str(OUTPUT)}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

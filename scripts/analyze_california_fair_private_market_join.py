#!/usr/bin/env python3
"""Join California FAIR share with the DOI voluntary-market ZIP file."""

from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "analysis/projects/us-housing-insurance-risk/data"
OUT = DATA / "california-fair-private-market-2022-join.json"


def read_fair() -> dict[str, dict[str, float]]:
    text = "\n".join(page.extract_text() or "" for page in PdfReader(DATA / "california-fair-plan-vs-voluntary-2022.pdf").pages)
    pattern = re.compile(r"^(.*?)\s+(\d{5})\s+(.*?)\s+([\d,]+|-|—)\s+([\d,]+|-|—)\s+([\d.]+)%\s*$")
    result = {}
    for match in map(pattern.match, text.splitlines()):
        if not match:
            continue
        county, zip_code, city, voluntary, fair, share = match.groups()
        result[zip_code] = {"fair_share": float(share), "fair_units": 0 if fair in ("-", "—") else int(fair.replace(",", "")), "voluntary_units": 0 if voluntary in ("-", "—") else int(voluntary.replace(",", ""))}
    return result


def main() -> int:
    fair = read_fair()
    ws = load_workbook(DATA / "california-residential-insurance-zip-2020-2023.xlsx", read_only=True, data_only=True).active
    statewide = defaultdict(lambda: {"new": 0, "renewed": 0, "nonrenewed": 0, "rows": 0})
    joined = []
    for county, zip_code, year, new, renewed, nonrenewed in ws.iter_rows(min_row=2, values_only=True):
        if year not in (2020, 2021, 2022, 2023):
            continue
        a = statewide[int(year)]
        a["rows"] += 1
        a["new"] += new or 0
        a["renewed"] += renewed or 0
        a["nonrenewed"] += nonrenewed or 0
        if year == 2022 and str(zip_code).zfill(5) in fair:
            f = fair[str(zip_code).zfill(5)]
            joined.append({"zip": str(zip_code).zfill(5), "fair_share": f["fair_share"], "renewed": renewed or 0, "nonrenewed": nonrenewed or 0})

    joined.sort(key=lambda x: x["fair_share"])
    q = len(joined) // 4
    quartiles = {}
    for label, values in [("lowest_fair_share_quartile", joined[:q]), ("highest_fair_share_quartile", joined[-q:])]:
        renew = sum(x["renewed"] for x in values)
        nonrenew = sum(x["nonrenewed"] for x in values)
        quartiles[label] = {"zip_rows": len(values), "mean_fair_share": round(sum(x["fair_share"] for x in values) / len(values), 6), "renewed": renew, "nonrenewed": nonrenew, "nonrenewal_rate_among_renewal_decisions": round(nonrenew / (renew + nonrenew), 6) if renew + nonrenew else None}
    result = {
        "format": "california-fair-private-market-join-v1",
        "period": "2020-2023 voluntary counts; 2022 FAIR share crosswalk",
        "voluntary_input": "california-residential-insurance-zip-2020-2023.xlsx",
        "fair_input": "california-fair-plan-vs-voluntary-2022.pdf",
        "statewide_voluntary_counts": {str(y): v for y, v in sorted(statewide.items())},
        "fair_zip_rows": len(fair),
        "joined_2022_zip_rows": len(joined),
        "join_rate_against_fair_rows": round(len(joined) / len(fair), 6) if fair else None,
        "quartile_comparison": quartiles,
        "classification": "ZIP-level descriptive crosswalk; FAIR share is residential dwelling structures while the voluntary file contains a broader set of residential policy forms and counts.",
        "limits": ["The two files do not share an identical denominator or policy universe.", "A ZIP-level association is not an individual household transition or causal private-market withdrawal estimate.", "CDI notes that 2020 onward nonrenewal/cancellation categories changed slightly from prior reports.", "Counts do not reveal reasons, premiums, coverage adequacy, claims, repairs, or moves."],
    }
    OUT.write_text(json.dumps(result, indent=2) + "\n")
    print(json.dumps({"fair_zip_rows": len(fair), "joined_2022_zip_rows": len(joined), "join_rate": result["join_rate_against_fair_rows"], "statewide": result["statewide_voluntary_counts"], "quartiles": quartiles}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

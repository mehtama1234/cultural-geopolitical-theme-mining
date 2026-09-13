#!/usr/bin/env python3
"""Summarize selected CBP sectors as shares of county establishments."""

from __future__ import annotations

import argparse
import csv
import io
import math
import statistics
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook


SECTORS = {
    "31----": "manufacturing",
    "44----": "retail",
    "62----": "health_social_assistance",
    "72----": "accommodation_food",
}


def pearson(xs: list[float], ys: list[float]) -> float:
    mean_x = sum(xs) / len(xs)
    mean_y = sum(ys) / len(ys)
    numerator = sum((x - mean_x) * (y - mean_y) for x, y in zip(xs, ys))
    denominator = math.sqrt(
        sum((x - mean_x) ** 2 for x in xs) * sum((y - mean_y) ** 2 for y in ys)
    )
    return numerator / denominator if denominator else float("nan")


def read_bfs(path: Path) -> dict[tuple[str, str], tuple[float, float, float]]:
    sheet = load_workbook(path, read_only=True, data_only=True).active
    values = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        try:
            key = (str(row[3]).zfill(2), str(row[4]).zfill(3))
            values[key] = tuple(float(str(row[i]).replace(",", "")) for i in (23, 24, 25))
        except (TypeError, ValueError, IndexError):
            continue
    return values


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--cbp", type=Path, required=True)
    parser.add_argument("--bfs", type=Path)
    parser.add_argument("--minimum-establishments", type=int, default=100)
    args = parser.parse_args()

    counties: dict[tuple[str, str], dict[str, int]] = {}
    with ZipFile(args.cbp) as archive:
        members = [name for name in archive.namelist() if name.endswith(".txt")]
        if len(members) != 1:
            raise ValueError(f"expected one CBP text member, found {members}")
        with archive.open(members[0]) as raw:
            reader = csv.DictReader(io.TextIOWrapper(raw, encoding="latin1"))
            for row in reader:
                if row.get("naics") not in ("------", *SECTORS) or not row.get("est", "").isdigit():
                    continue
                key = (row["fipstate"], row["fipscty"])
                counties.setdefault(key, {})[row["naics"]] = int(row["est"])

    eligible = [
        values for values in counties.values()
        if values.get("------", 0) >= args.minimum_establishments
    ]
    print(f"counties={len(eligible)}")
    for code, label in SECTORS.items():
        shares = [100 * values.get(code, 0) / values["------"] for values in eligible]
        quartiles = statistics.quantiles(shares, n=4, method="inclusive")
        print(
            f"{label}\tmedian={statistics.median(shares):.2f}%"
            f"\tp25={quartiles[0]:.2f}%\tp75={quartiles[2]:.2f}%"
        )

    if args.bfs:
        bfs = read_bfs(args.bfs)
        matched = [
            (bfs[key], values)
            for key, values in counties.items()
            if key in bfs and values.get("------", 0) >= args.minimum_establishments
        ]
        for year_index, year in enumerate((2023, 2024, 2025)):
            applications_per_establishment = [
                application[year_index] / values["------"] for application, values in matched
            ]
            for code, label in SECTORS.items():
                sector_share = [values.get(code, 0) / values["------"] for _, values in matched]
                print(
                    f"pearson_ba{year}_over_est_vs_{label}="
                    f"{pearson(applications_per_establishment, sector_share):.3f}"
                )


if __name__ == "__main__":
    main()
